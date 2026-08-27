from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT=Path(r"C:\Users\kccistc\Desktop\kernel-cc1101-spi\output\fhss-ota-radio-presentation-script.docx")
ORANGE=RGBColor(0xC8,0x43,0x08); INK=RGBColor(0x17,0x14,0x12); GRAY=RGBColor(0x6F,0x6A,0x65)
FONT="Malgun Gothic"

slides=[
("1. 주파수 호핑 기반 무전기 OTA 시스템","약 50초","안녕하세요. 저희는 라즈베리파이 게이트웨이에서 여러 ESP32 무전기로 펌웨어를 무선 배포하고, 같은 장치가 음성 통신까지 수행하도록 만든 FHSS OTA Radio 프로젝트를 소개하겠습니다. 발표에서는 전체 구조를 먼저 보고, 드라이버와 Yocto, Gateway와 Qt, 공통 프로토콜, FHSS, 음성 펌웨어, ESP32 OTA 순서로 설명하겠습니다. 마지막에는 실제 통합 과정에서 어떤 문제가 있었고 어디까지 검증했는지 말씀드리겠습니다.","첫 문장은 천천히. 프로젝트를 ‘무선 업데이트 + 무전 음성’으로 한 문장에 정의합니다."),
("2. 프로젝트 목표","약 55초","프로젝트의 목표는 단순히 파일을 무선으로 보내는 것이 아닙니다. 첫째, Gateway가 여러 무전기를 찾고 업데이트를 관리해야 합니다. 둘째, 손실이 있는 무선 링크에서도 ACK와 NACK로 필요한 패킷만 다시 보내야 합니다. 셋째, 간섭을 피하기 위해 Gateway와 무전기가 같은 채널 순서로 이동해야 합니다. 마지막으로 ESP32는 받은 파일의 해시를 검증한 뒤 안전한 OTA 파티션으로 부팅해야 합니다. 핵심은 서로 다른 계층이 하나의 규칙과 시간 기준으로 움직이게 만드는 것입니다.","네 목표를 모두 읽기보다 마지막 핵심 문장을 강조합니다."),
("3. 전체 시스템 구조","약 65초","전체 구조를 보면 왼쪽의 Gateway 앱이 업데이트 파일과 대상 장치를 관리합니다. 앱은 /dev/cc1101에 기본적인 read, write, ioctl을 요청하고, Linux 드라이버가 SPI 레지스터, FIFO, 인터럽트와 FHSS 시간을 담당합니다. 무선 구간에서는 데이터 패킷뿐 아니라 ACK와 동기 패킷이 오갑니다. ESP32 펌웨어는 OTA 데이터를 플래시에 쓰는 동시에 음성 코덱과 사용자 인터페이스도 담당합니다. 공통 프로토콜은 제어 규칙을, FHSS는 시간과 채널 규칙을 연결합니다.","왼쪽에서 오른쪽으로 손이나 포인터를 이동하며 설명합니다."),
("4. Device Driver + Yocto","약 75초","드라이버 파트의 설계 원칙은 앱을 단순하게 만드는 것입니다. 앱은 일반 파일처럼 장치를 열고 read, write, ioctl, poll만 사용합니다. 커널 내부에서 SPI 레지스터 접근, RX와 TX FIFO, GDO 인터럽트, 채널 전환과 SYNC 송신을 처리합니다. RX FIFO overflow에서는 IDLE, SFRX, SRX 순서로 자동 복구하고, 사용자 데이터와 내부 SYNC가 동시에 송신되지 않도록 tx_lock으로 직렬화했습니다. Yocto에는 모듈과 Device Tree overlay, 시험 도구를 recipe로 묶어 같은 이미지를 반복 빌드할 수 있게 했습니다. 대상 커널은 Linux 5.15.92입니다.","내 담당 파트라면 RX 복구와 tx_lock의 이유를 한 문장씩 분명히 말합니다."),
("5. Gateway + Qt","약 70초","Gateway에서는 사용자가 보는 Qt 화면과 실제 전송 상태 기계를 분리했습니다. 사용자는 파일을 선택하고 장치를 탐색한 뒤 진행률과 로그를 확인합니다. 내부 OtaSession은 START로 세션을 열고, 데이터를 배치로 보내며, ACK와 NACK를 보고 빠진 sequence만 재전송한 뒤 END와 해시 검증으로 종료합니다. 고정 채널에서는 전체 파일 전송과 SHA-256 일치를 확인했습니다. FHSS에서는 패킷을 슬롯 경계와 guard time에 맞추는 작업이 남아 있어, 기능은 통합됐지만 처리량은 최적화 중입니다.","완료와 최적화 중인 항목을 섞지 않습니다."),
("6. 공통 프로토콜","약 60초","ota-protocol 저장소에는 탐색, OTA, FHSS 설정과 동기 패킷의 wire format을 정의했습니다. FHSS 설정에는 generation, seed, 채널 수, 슬롯 시간, 알고리즘 버전이 들어갑니다. seed는 실제 채널 목록이 아니라 양쪽이 같은 순서를 계산하기 위한 입력입니다. Golden test vector를 두어 같은 입력이 같은 바이트열과 같은 채널 순서를 만드는지 확인하고, 알고리즘 버전이 다르면 명확하게 거부하도록 했습니다.","seed의 의미를 ‘채널표가 아닌 순서 생성 입력’이라고 쉽게 설명합니다."),
("7. FHSS 동작 원리","약 65초","FHSS에서는 seed로 xorshift32 기반 permutation을 만들고, 현재 슬롯 번호를 이용해 채널을 선택합니다. 예를 들어 확인한 순서는 1, 7, 6, 5, 3, 2, 8, 4였습니다. 한 슬롯은 300밀리초이고, 슬롯 시작 전후 5밀리초는 채널을 바꾸는 guard time입니다. 채널 1은 최초 접속과 동기 상실 때 다시 만나는 rendezvous 채널입니다. xorshift32는 보안을 위한 암호 난수가 아니라, 같은 seed와 버전으로 같은 순서를 재현하기 위한 결정적 알고리즘입니다.","슬롯은 ‘한 채널에 머무는 시간 구간’이라고 덧붙이면 이해가 쉽습니다."),
("8. 동기화와 복구","약 70초","Gateway가 CONFIG를 보내면 ESP32는 seed와 generation을 대기 설정으로 저장합니다. ACTIVATE를 받으면 호핑 준비 상태로 전환하고, 채널 1에서 MASTER의 SYNC를 받아 현재 슬롯과 시간 기준을 계산합니다. TRACKING 상태에서는 매 슬롯의 SYNC로 오차를 보정합니다. SYNC를 연속으로 놓치면 채널 1로 돌아와 같은 세션을 유지한 채 재동기합니다. 실기기 로그에서는 SYNC_ACQUIRED에서 TRACKING으로 이동했고 시간 오차는 약 마이너스 22에서 플러스 18마이크로초 범위였습니다.","CONFIG→ACTIVATE→SYNC 순서를 또렷하게 읽습니다."),
("9. 음성 압축 + 펌웨어","약 65초","ESP32의 음성 경로는 마이크 입력, 압축, 무선 전송, 수신 큐, 스피커 출력으로 이어집니다. INMP441에서 8킬로헤르츠 PCM을 받고 Speex narrowband로 압축해 두 프레임 단위 패킷으로 보냅니다. 수신 쪽은 큐로 순간 지연을 흡수하고 END 제어와 silence tail을 처리한 뒤 MAX98357A로 재생합니다. 사용자는 PTT와 회전 인코더를 조작하고 OLED와 LED로 상태를 확인합니다. 평상시에는 무전기이고 업데이트 모드에서는 OTA client가 됩니다.","오디오 파이프라인을 왼쪽부터 한 번에 연결해 설명합니다."),
("10. ESP32 OTA","약 65초","ESP32 OTA는 실행 중인 파티션을 바로 덮어쓰지 않습니다. Discovery에서 장치 ID와 버전을 알리고, START에서 파일 크기와 SHA-256, session을 확인합니다. DATA는 sequence를 검사하면서 비활성 OTA 파티션에 쓰고, 배치 ACK와 NACK로 누락 정보를 보냅니다. END에서 전체 SHA-256이 일치해야 다음 부팅 파티션을 변경합니다. FHSS 설정도 CONFIG에서는 pending으로 저장하고 ACTIVATE 이후 active로 전환합니다.","오른쪽 dual partition 그림에서 안전성을 강조합니다."),
("11. 통합 트러블슈팅","약 90초","통합 과정에서 수신이 멎는 문제는 RF 감도가 아니라 RX FIFO overflow 복구 누락이었습니다. 슬롯 1부터 SYNC가 거부된 문제는 hop_index 의미가 서로 달랐기 때문입니다. 초기 획득 실패는 같은 슬롯의 SYNC를 짧게 반복한 문제였고, OTA_DATA write 실패는 내부 SYNC와 사용자 DATA가 TX 경로에서 충돌했기 때문입니다. 각각 상태 복구, 의미 통일, 슬롯당 한 번의 SYNC, tx_lock으로 해결했습니다. 현재 남은 성능 문제는 300밀리초 슬롯당 한 패킷만 보내는 보수적인 정책입니다. 패킷 구조뿐 아니라 state, index, slot 경계의 의미까지 저장소 사이에서 같아야 한다는 점을 배웠습니다.","가장 중요한 기술 슬라이드입니다. 현상→원인→조치 순서를 유지합니다."),
("12. 검증 결과와 현재 상태","약 70초","고정 채널 OTA에서는 7,829개 패킷 전체 수신과 SHA-256 일치를 확인했습니다. FHSS 제어에서는 CONFIG와 ACTIVATE 대상 ACK를 확인했고, ESP32가 SYNC_ACQUIRED에서 TRACKING으로 이동했습니다. 시간 오차는 약 마이너스 22에서 플러스 18마이크로초였고 RSSI와 CRC도 정상 범위였습니다. 다만 호핑 중 대용량 파일 전송은 현재 슬롯당 전송량이 보수적이어서 처리량 최적화가 필요합니다. 구현됨, 실기기 검증됨, 최적화 중인 항목을 구분해 평가했습니다.","수치는 자신 있게 말하되 FHSS OTA 완전 성공으로 과장하지 않습니다."),
("13. 결론과 다음 단계","약 60초","정리하면 공통 프로토콜로 저장소 사이의 계약을 만들고, Linux 드라이버가 채널 전환과 동기화처럼 시간에 민감한 일을 맡도록 했습니다. Gateway는 재전송 상태 기계와 사용자 인터페이스를, ESP32는 dual-partition OTA와 음성 기능을 담당합니다. 그 결과 CONFIG와 ACTIVATE, FHSS 동기 추적, 고정 채널 OTA 전체 검증까지 연결했습니다. 다음 단계는 한 슬롯 안에서 여러 데이터 패킷과 ACK를 안전하게 처리해 속도를 높이고, 장시간 재동기와 여러 무전기 동시 업데이트를 시험하는 것입니다. 이상으로 발표를 마치겠습니다.","마지막 문장은 화면을 보지 말고 청중을 보며 마무리합니다."),
]

def set_font(run,size=None,bold=None,color=None):
    run.font.name=FONT; run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),FONT)
    if size: run.font.size=Pt(size)
    if bold is not None: run.bold=bold
    if color: run.font.color.rgb=color

doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.8); sec.right_margin=Inches(.8)
styles=doc.styles
styles['Normal'].font.name=FONT; styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'),FONT); styles['Normal'].font.size=Pt(10.5)
for st,size,color in [('Title',28,INK),('Heading 1',18,INK),('Heading 2',12,ORANGE)]:
    styles[st].font.name=FONT; styles[st]._element.rPr.rFonts.set(qn('w:eastAsia'),FONT); styles[st].font.size=Pt(size); styles[st].font.color.rgb=color

p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.LEFT; p.paragraph_format.space_before=Pt(55)
r=p.add_run("FHSS-OTA RADIO"); set_font(r,12,True,ORANGE)
p=doc.add_paragraph(); p.style='Title'; p.paragraph_format.space_before=Pt(12); p.paragraph_format.space_after=Pt(18); p.add_run("최종 발표 스크립트")
p=doc.add_paragraph("Device Driver + Yocto · OTA + Qt · FHSS · Audio + Firmware · ESP32 + OTA"); p.paragraph_format.space_after=Pt(30)
for run in p.runs:set_font(run,13,False,GRAY)
p=doc.add_paragraph(); r=p.add_run("발표 길이  약 15분  |  총 13장"); set_font(r,12,True,INK)
p=doc.add_paragraph("사용법: 굵은 슬라이드 제목으로 흐름을 익힌 뒤, 본문은 문장 그대로 암기하기보다 핵심어를 중심으로 자연스럽게 말합니다.")
for run in p.runs:set_font(run,10.5,False,GRAY)
doc.add_page_break()

for idx,(title,timing,script,cue) in enumerate(slides):
    h=doc.add_paragraph(style='Heading 1'); h.paragraph_format.space_after=Pt(3); h.add_run(title)
    t=doc.add_paragraph(style='Heading 2'); t.paragraph_format.space_after=Pt(8); t.add_run(timing)
    p=doc.add_paragraph(script); p.paragraph_format.line_spacing=1.35; p.paragraph_format.space_after=Pt(10)
    for run in p.runs:set_font(run,10.5,False,INK)
    q=doc.add_paragraph(); q.paragraph_format.left_indent=Inches(.18); q.paragraph_format.space_after=Pt(16)
    r=q.add_run("발표 포인트  "); set_font(r,10,True,ORANGE); r=q.add_run(cue); set_font(r,10,False,GRAY)
    if idx in [2,5,8,10]: doc.add_page_break()

for section in doc.sections:
    footer=section.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=footer.add_run("FHSS-OTA Radio · Presentation Script"); set_font(r,8,False,GRAY)
OUT.parent.mkdir(parents=True,exist_ok=True); doc.save(OUT)
print(OUT)
